import pytest
from pathlib import Path
from docx import Document
from src.process_docs import replace_text_in_docx

# Globals to match your function
FIND_TEXT = "World"
REPLACE_TEXT = "Universe"


@pytest.fixture
def tmp_docx(tmp_path):
    """Create a temporary .docx file for testing."""
    def _make_docx(text: str, filename: str = "test.docx") -> Path:
        file_path = tmp_path / filename
        doc = Document()
        doc.add_paragraph(text)
        doc.save(file_path)
        return file_path
    return _make_docx


def get_docx_text(path: Path) -> str:
    doc = Document(path)
    return "\n".join([p.text for p in doc.paragraphs])


def test_replacement(tmp_docx):
    infile = tmp_docx("Hello World!")
    outfile = infile.parent / "out.docx"

    replace_text_in_docx(infile, outfile, FIND_TEXT, REPLACE_TEXT)
    assert REPLACE_TEXT in get_docx_text(outfile)


def test_no_match(tmp_docx):
    infile = tmp_docx("Hello again")
    outfile = infile.parent / "out.docx"

    replace_text_in_docx(infile, outfile, ["Missing"], "Replaced")
    assert "Replaced" not in get_docx_text(outfile)
    assert get_docx_text(infile) == get_docx_text(outfile)


def test_empty_file(tmp_docx):
    infile = tmp_docx("")
    outfile = infile.parent / "out.docx"

    replace_text_in_docx(infile, outfile, "Anything", "Replacement")
    assert get_docx_text(outfile) == ""
